"""Erstellt Plattenwechsler-Dokumentation als Word-Datei."""
import io, math, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
import matplotlib.patheffects as pe
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Farben ──────────────────────────────────────────────────────────────
C_RED   = "#C8102E"
C_DARK  = "#1E2028"
C_MED   = "#2D3142"
C_LIGHT = "#6C7280"
C_OK    = "#22C55E"
C_WARN  = "#F59E0B"
C_INFO  = "#3B82F6"
C_BG    = "#F8F9FA"
C_WHITE = "#FFFFFF"

# ── Hilfsfunktion: Bild als BytesIO ─────────────────────────────────────
def fig_to_stream(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight",
                facecolor=fig.get_facecolor())
    buf.seek(0)
    plt.close(fig)
    return buf

# ════════════════════════════════════════════════════════════════════════
# Grafik 1: Systemarchitektur
# ════════════════════════════════════════════════════════════════════════
def make_architektur():
    fig, ax = plt.subplots(figsize=(13, 7.5))
    fig.patch.set_facecolor(C_BG)
    ax.set_facecolor(C_BG)
    ax.set_xlim(0, 13); ax.set_ylim(0, 7.5)
    ax.axis("off")

    def box(x, y, w, h, label, sublabel="", color=C_MED, fontsize=10, textcolor=C_WHITE):
        rect = FancyBboxPatch((x, y), w, h,
                              boxstyle="round,pad=0.08", linewidth=1.5,
                              edgecolor=color, facecolor=color)
        ax.add_patch(rect)
        ax.text(x+w/2, y+h/2+(0.15 if sublabel else 0), label,
                ha="center", va="center", color=textcolor,
                fontsize=fontsize, fontweight="bold")
        if sublabel:
            ax.text(x+w/2, y+h/2-0.25, sublabel,
                    ha="center", va="center", color=C_BG,
                    fontsize=7.5, style="italic")

    def arrow(x1, y1, x2, y2, label="", color=C_LIGHT, lw=1.5):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="-|>", color=color,
                                   lw=lw, mutation_scale=14))
        if label:
            mx, my = (x1+x2)/2, (y1+y2)/2
            ax.text(mx+0.08, my, label, fontsize=7, color=C_LIGHT)

    def darrow(x1, y1, x2, y2, label="", color=C_LIGHT):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="<|-|>", color=color,
                                   lw=1.5, mutation_scale=13))
        if label:
            mx, my = (x1+x2)/2, (y1+y2)/2
            ax.text(mx+0.08, my, label, fontsize=7, color=C_LIGHT)

    # ── Raspberry Pi Hauptblock ──────────────────────────────────────
    pi_rect = FancyBboxPatch((0.3, 0.4), 8.2, 6.8,
                             boxstyle="round,pad=0.15", linewidth=2.5,
                             edgecolor=C_RED, facecolor="#FFF5F5", zorder=0)
    ax.add_patch(pi_rect)
    ax.text(4.4, 7.0, "Raspberry Pi 4", ha="center", va="center",
            color=C_RED, fontsize=12, fontweight="bold")

    # Module
    box(0.6, 5.3, 2.4, 0.9, "Hauptablauf", "Statemachine", C_RED, 9)
    box(0.6, 4.1, 2.4, 0.9, "ESP-Client", "UART-Treiber", C_INFO, 9)
    box(0.6, 2.9, 2.4, 0.9, "GPIO-Manager", "Endschalter / Pins", C_MED, 9)
    box(0.6, 1.7, 2.4, 0.9, "MQTT-Client", "Broker-Verbindung", "#7C3AED", 9)
    box(0.6, 0.6, 2.4, 0.9, "Telegram-Client", "Bot-API", "#0EA5E9", 9)
    box(3.6, 4.1, 2.0, 0.9, "AuftragsQueue", "Thread-safe FIFO", "#64748B", 9)
    box(3.6, 2.9, 2.0, 0.9, "FehlerBehandlung", "Fehler & Quittierung", "#DC2626", 9)
    box(3.6, 1.7, 2.0, 0.9, "Config", "YAML lesen/schreiben", "#059669", 9)
    box(5.9, 5.3, 2.4, 0.9, "HMI / PyQt5", "Touchscreen-UI", "#D97706", 9)
    box(5.9, 4.1, 2.4, 0.9, "main.py", "Startpunkt / DI", "#475569", 9)

    # Interne Pfeile
    arrow(1.8, 5.3, 1.8, 5.0)
    arrow(1.8, 4.1, 1.8, 3.8)
    arrow(1.8, 2.9, 1.8, 2.6)
    arrow(3.0, 5.75, 3.6, 5.75, "", C_RED)
    arrow(3.0, 4.55, 3.6, 4.55)
    arrow(3.0, 3.35, 3.6, 3.35)
    arrow(3.0, 2.15, 3.6, 2.15)
    arrow(5.6, 4.55, 5.9, 4.55)
    arrow(7.1, 4.1, 7.1, 6.2)
    arrow(5.9, 5.75, 3.0, 5.75, "", C_RED)

    # ── Externe Blöcke ───────────────────────────────────────────────
    box(9.2, 5.4, 3.3, 1.2, "ESP32", "UART 115200 Baud", C_MED, 10)
    box(9.2, 3.9, 3.3, 1.2, "MQTT-Broker", "HiveMQ / lokal", "#7C3AED", 10)
    box(9.2, 2.4, 3.3, 1.2, "Telegram", "Bot-Server", "#0EA5E9", 10)
    box(9.2, 0.8, 1.3, 1.2, "GPIO\nEndschalter", "", C_DARK, 9)
    box(10.8, 0.8, 1.7, 1.2, "GPIO\nDrucker-Pins", "", C_DARK, 9)

    # Externe Pfeile
    darrow(8.5, 4.55, 9.2, 6.0, "UART")
    darrow(8.5, 2.15, 9.2, 4.5, "TCP/IP")
    darrow(8.5, 0.75+0.6, 9.2, 2.75+0.3, "HTTPS")
    arrow(9.2+0.65, 2.0, 1.8, 2.9, "GPIO-IRQ", "#374151")
    arrow(11.65, 2.0, 1.8+0.3, 1.7+0.9)

    ax.text(4.4, -0.05, "Interne Module", ha="center", fontsize=8,
            color=C_LIGHT, style="italic")
    return fig_to_stream(fig)

# ════════════════════════════════════════════════════════════════════════
# Grafik 2: Statemachine
# ════════════════════════════════════════════════════════════════════════
def make_statemachine():
    fig, ax = plt.subplots(figsize=(13, 8))
    fig.patch.set_facecolor(C_BG)
    ax.set_facecolor(C_BG)
    ax.set_xlim(0, 13); ax.set_ylim(0, 8)
    ax.axis("off")

    states = {
        "INIT":          (2.0, 6.8, C_MED),
        "REFERENZFAHRT": (2.0, 5.3, C_INFO),
        "BEREITSCHAFT":  (2.0, 3.8, C_OK),
        "PLATTENWECHSEL":(6.5, 3.8, C_RED),
        "SERVICE":       (6.5, 5.3, "#7C3AED"),
        "FEHLER":        (6.5, 6.8, "#DC2626"),
        "NOT_AUS":       (10.5, 6.8, "#991B1B"),
    }

    def state_box(x, y, label, color):
        w, h = 2.8, 0.9
        rect = FancyBboxPatch((x-w/2, y-h/2), w, h,
                              boxstyle="round,pad=0.1", linewidth=2,
                              edgecolor=color, facecolor=color)
        ax.add_patch(rect)
        ax.text(x, y, label, ha="center", va="center",
                color=C_WHITE, fontsize=9, fontweight="bold")

    for name, (x, y, c) in states.items():
        state_box(x, y, name, c)

    def arr(src, dst, label="", rad=0.0, lbl_off=(0,0)):
        x1, y1, _ = states[src]
        x2, y2, _ = states[dst]
        style = f"arc3,rad={rad}"
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="-|>", color=C_LIGHT,
                                   lw=1.4, mutation_scale=13,
                                   connectionstyle=style))
        if label:
            mx = (x1+x2)/2 + lbl_off[0]
            my = (y1+y2)/2 + lbl_off[1]
            ax.text(mx, my, label, fontsize=7, color=C_DARK,
                    ha="center",
                    bbox=dict(boxstyle="round,pad=0.15", fc=C_BG,
                              ec="none", alpha=0.85))

    arr("INIT",          "REFERENZFAHRT", "ESP verbunden",        lbl_off=(0.5,0))
    arr("REFERENZFAHRT", "BEREITSCHAFT",  "HOME_DONE",            lbl_off=(0.5,0))
    arr("BEREITSCHAFT",  "PLATTENWECHSEL","Auftrag in Queue",     lbl_off=(0,0.2))
    arr("PLATTENWECHSEL","BEREITSCHAFT",  "Wechsel fertig",       lbl_off=(0,-0.25))
    arr("BEREITSCHAFT",  "SERVICE",       "Service aktivieren",   rad=-0.3, lbl_off=(-0.2,0.6))
    arr("SERVICE",       "BEREITSCHAFT",  "Service verlassen",    rad=-0.3, lbl_off=(1.2,-0.6))
    arr("REFERENZFAHRT", "FEHLER",        "Timeout / Fehler",     lbl_off=(1.8,0.2))
    arr("PLATTENWECHSEL","FEHLER",        "Fehler",               lbl_off=(0.3,0.5))
    arr("BEREITSCHAFT",  "FEHLER",        "ESP-Fehler",           rad=-0.2, lbl_off=(2.5,0.2))
    arr("FEHLER",        "REFERENZFAHRT", "Quittiert\n+ Reset",  rad=0.35, lbl_off=(-1.5,0))
    arr("FEHLER",        "NOT_AUS",       "Not-Aus",              lbl_off=(0,0.25))
    arr("NOT_AUS",       "FEHLER",        "Quittiert",            lbl_off=(0,-0.25))

    # REFERENZFAHRT Schleife (manuelle Referenzfahrt aus BEREITSCHAFT)
    ax.annotate("", xy=(2.0, 4.25), xytext=(2.0, 4.9),
                arrowprops=dict(arrowstyle="-|>", color=C_LIGHT,
                               lw=1.2, mutation_scale=11,
                               connectionstyle="arc3,rad=0.6"))
    ax.text(0.3, 4.55, "Manuelle\nReferenzfahrt", fontsize=6.5,
            color=C_DARK, ha="center")

    ax.set_title("Systemzustände — Raspberry Pi", fontsize=13,
                 fontweight="bold", color=C_DARK, pad=10)
    return fig_to_stream(fig)

# ════════════════════════════════════════════════════════════════════════
# Grafik 3: Plattenwechsel-Sequenz
# ════════════════════════════════════════════════════════════════════════
def make_sequenz():
    fig, ax = plt.subplots(figsize=(13, 9.5))
    fig.patch.set_facecolor(C_BG)
    ax.set_facecolor(C_BG)
    ax.set_xlim(0, 13); ax.set_ylim(0, 9.5)
    ax.axis("off")

    phases = [
        ("Phase 1 — Platte aus Drucker holen", C_RED, [
            "1. MOVE_TO (pos_x, pos_z_anfahr)  →  Ausgangsposition",
            "2. OPEN_DOOR  →  ESP fährt intern zur Tür, öffnet Kreisbogen, kehrt zurück",
            "3. MOVE_TO (pos_x, pos_z_druckbett)  →  Gabel-Bereitschaftsposition",
            "4. PICKUP (gripper_depth, lift_offset)  →  Platte aufnehmen",
            "5. MOVE_TO (pos_x, pos_z_anfahr)  →  zurück Ausgangsposition",
            "6. CLOSE_DOOR  →  ESP schließt Tür per Kreisbogen, kehrt zurück",
        ]),
        ("Phase 2 — Alte Platte ablegen", "#D97706", [
            "7. MOVE_TO (ablage.x, ablage.z)  →  Ablageposition anfahren",
            "8. DEPOSIT (gripper_depth, lift_offset)  →  Platte ablegen",
            "   → Ablage wird in Config als BELEGT markiert",
        ]),
        ("Phase 3 — Neue Platte aus Magazin holen", "#059669", [
            "9.  MOVE_TO (magazin.x, magazin.z)  →  Magazinposition anfahren",
            "10. PICKUP (gripper_depth, lift_offset)  →  Neue Platte aufnehmen",
            "    → Magazin wird in Config als LEER markiert",
        ]),
        ("Phase 4 — Neue Platte einlegen", C_INFO, [
            "11. MOVE_TO (pos_x, pos_z_anfahr)  →  Ausgangsposition",
            "12. OPEN_DOOR  →  Tür öffnen",
            "13. MOVE_TO (pos_x, pos_z_druckbett)  →  Einlegeposition",
            "14. DEPOSIT (gripper_depth, lift_offset)  →  Platte einlegen",
            "15. MOVE_TO (pos_x, pos_z_anfahr)  →  zurück Ausgangsposition",
            "16. CLOSE_DOOR  →  Tür schließen",
        ]),
        ("Phase 5 — Heimfahrt", C_LIGHT, [
            "17. MOVE_HOME  →  nur wenn Queue leer (sonst nächster Auftrag direkt)",
        ]),
    ]

    y = 9.1
    for title, color, steps in phases:
        # Phasentitel
        rect = FancyBboxPatch((0.3, y-0.38), 12.4, 0.45,
                              boxstyle="round,pad=0.05",
                              facecolor=color, edgecolor="none")
        ax.add_patch(rect)
        ax.text(0.6, y-0.15, title, color=C_WHITE,
                fontsize=9.5, fontweight="bold", va="center")
        y -= 0.55

        for step in steps:
            indent = 0.3 if step.startswith("   ") else 0.0
            ax.text(0.7+indent, y, step, fontsize=8.2, color=C_DARK,
                    va="center", family="monospace")
            y -= 0.36
        y -= 0.15

    # Sonderfall-Box
    ax.add_patch(FancyBboxPatch((0.3, 0.05), 12.4, 0.5,
                                boxstyle="round,pad=0.05",
                                facecolor="#FEF3C7", edgecolor="#F59E0B", lw=1.5))
    ax.text(0.6, 0.3,
            "⚠  Fehler während Plattenwechsel: Queue wird geleert — "
            "Ablage, Magazin und Schlitten manuell prüfen!",
            fontsize=8, color="#92400E", va="center")

    ax.set_title("Plattenwechsel — Ablaufsequenz", fontsize=13,
                 fontweight="bold", color=C_DARK, pad=8)
    return fig_to_stream(fig)

# ════════════════════════════════════════════════════════════════════════
# Grafik 4: Kommunikationsprotokoll
# ════════════════════════════════════════════════════════════════════════
def make_protokoll():
    fig, ax = plt.subplots(figsize=(13, 7))
    fig.patch.set_facecolor(C_BG)
    ax.set_facecolor(C_BG)
    ax.set_xlim(0, 13); ax.set_ylim(0, 7)
    ax.axis("off")

    # Lebenslinen
    for x, lbl, col in [(2.5,"Raspberry Pi", C_RED),(10.5,"ESP32", C_MED)]:
        ax.plot([x,x],[0.3,6.5], color=col, lw=1.5, ls="--", alpha=0.4)
        rect = FancyBboxPatch((x-1.1, 6.3), 2.2, 0.6,
                              boxstyle="round,pad=0.08",
                              facecolor=col, edgecolor="none")
        ax.add_patch(rect)
        ax.text(x, 6.6, lbl, ha="center", va="center",
                color=C_WHITE, fontsize=10, fontweight="bold")

    msgs = [
        # (y,  dir,   label,              note,                  col)
        (5.9, "→", "CMD;1;HOME",          "Referenzfahrt starten", C_INFO),
        (5.5, "←", "RSP;1;ACK",           "Kommando akzeptiert",   C_OK),
        (5.1, "←", "EVT;0;STATE;BUSY_HOMING", "Zustand wechselt", C_WARN),
        (4.7, "→", "CMD;2;HOME_SWITCH_HIT;axis=X", "Pi meldet Endschalter X", C_INFO),
        (4.3, "→", "CMD;3;HOME_SWITCH_HIT;axis=Z", "Pi meldet Endschalter Z", C_INFO),
        (3.9, "←", "EVT;1;OK;HOME_DONE", "Referenzfahrt fertig",  C_OK),
        (3.5, "←", "EVT;0;STATE;READY",  "System bereit",         C_OK),
        (3.0, "→", "CMD;4;OPEN_DOOR;x_approach=…;z_approach=…;…", "Tür öffnen", C_RED),
        (2.6, "←", "RSP;4;ACK",          "",                      C_OK),
        (2.2, "←", "EVT;4;OK;DOOR_OPEN_DONE", "Tür geöffnet",    C_OK),
        (1.8, "→", "CMD;5;PICKUP;gripper_depth=120;lift_offset=8", "Platte holen", C_INFO),
        (1.4, "←", "EVT;5;OK;PICKUP_DONE","Platte aufgenommen",   C_OK),
        (0.9, "←", "EVT;0;HEARTBEAT;uptime_ms=…", "alle 1000 ms", C_LIGHT),
    ]

    for y, direction, label, note, col in msgs:
        if direction == "→":
            x1, x2 = 2.5, 10.5
        else:
            x1, x2 = 10.5, 2.5
        ax.annotate("", xy=(x2, y), xytext=(x1, y),
                    arrowprops=dict(arrowstyle="-|>", color=col,
                                   lw=1.5, mutation_scale=12))
        mid = (x1+x2)/2
        bg = "#EFF6FF" if direction == "→" else "#F0FDF4"
        ax.text(mid, y+0.12, label, ha="center", va="bottom",
                fontsize=7.2, color=C_DARK, family="monospace",
                bbox=dict(boxstyle="round,pad=0.15", fc=bg, ec="none"))
        if note:
            side_x = 11.3 if direction == "←" else 0.1
            ha = "left" if direction == "←" else "left"
            ax.text(side_x if direction=="←" else 0.1, y+0.08, note,
                    ha="left", va="center", fontsize=6.8,
                    color=C_LIGHT, style="italic")

    ax.set_title("Kommunikationsprotokoll Pi ↔ ESP32 (Auszug)",
                 fontsize=12, fontweight="bold", color=C_DARK, pad=8)
    return fig_to_stream(fig)

# ════════════════════════════════════════════════════════════════════════
# Grafik 5: Lager-Konzept (Ablagen + Magazine)
# ════════════════════════════════════════════════════════════════════════
def make_lager():
    fig, ax = plt.subplots(figsize=(11, 4.5))
    fig.patch.set_facecolor(C_BG)
    ax.set_facecolor(C_BG)
    ax.set_xlim(0, 11); ax.set_ylim(0, 4.5)
    ax.axis("off")

    def slot(x, y, label, status, color):
        rect = FancyBboxPatch((x, y), 2.2, 1.4,
                              boxstyle="round,pad=0.1",
                              facecolor=color, edgecolor=C_WHITE, lw=1.5)
        ax.add_patch(rect)
        ax.text(x+1.1, y+0.95, label, ha="center", va="center",
                color=C_WHITE, fontsize=9, fontweight="bold")
        ax.text(x+1.1, y+0.45, status, ha="center", va="center",
                color=C_WHITE, fontsize=8)

    # Ablagen
    ax.text(2.5, 4.2, "Ablagen", ha="center", fontsize=11,
            fontweight="bold", color=C_DARK)
    slot(0.3, 2.5, "Ablage 1", "● FREI",   C_OK)
    slot(2.7, 2.5, "Ablage 2", "● BELEGT", "#DC2626")
    slot(5.1, 2.5, "Ablage 3", "● FREI",   C_OK)

    # Magazine
    ax.text(8.5, 4.2, "Magazine", ha="center", fontsize=11,
            fontweight="bold", color=C_DARK)
    slot(6.8, 2.5, "Magazin 1", "● VERFÜGBAR", C_OK)
    slot(8.8+0.2, 2.5, "Magazin 2", "● LEER",  C_LIGHT)

    # Logik-Pfeile
    ax.text(5.5, 1.8, "Automatische Auswahl:", ha="center",
            fontsize=9, color=C_DARK, fontweight="bold")
    ax.text(5.5, 1.35,
            "naechste_freie_ablage()  →  erste Ablage mit Status FREI",
            ha="center", fontsize=8.5, color=C_DARK)
    ax.text(5.5, 0.95,
            "naechstes_verfuegbares_magazin()  →  erstes Magazin mit Status VERFÜGBAR",
            ha="center", fontsize=8.5, color=C_DARK)
    ax.text(5.5, 0.45,
            "Status wird nach DEPOSIT / PICKUP automatisch in config.yaml gespeichert",
            ha="center", fontsize=8, color=C_LIGHT, style="italic")

    ax.set_title("Lagerkonzept — Ablagen & Magazine", fontsize=12,
                 fontweight="bold", color=C_DARK, pad=8)
    return fig_to_stream(fig)

# ════════════════════════════════════════════════════════════════════════
# Word-Dokument
# ════════════════════════════════════════════════════════════════════════
def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E) if level == 1 else RGBColor(0x1E, 0x20, 0x28)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    h.paragraph_format.space_after  = Pt(6)
    return h

def add_para(doc, text, bold=False, italic=False, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_mono(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(2)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = "Table Grid"
    # Header
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
        # Hintergrundfarbe Header
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "C8102E")
        shading.set(qn("w:color"), "FFFFFF")
        shading.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Daten
    for ri, row in enumerate(rows):
        trow = table.rows[ri+1]
        for ci, val in enumerate(row):
            cell = trow.cells[ci]
            cell.text = str(val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
        if ri % 2 == 0:
            for ci in range(len(headers)):
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F8F9FA")
                shading.set(qn("w:val"), "clear")
                trow.cells[ci]._tc.get_or_add_tcPr().append(shading)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table

def add_img(doc, stream, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(stream, width=Cm(width_cm))
    doc.add_paragraph()

def add_info_box(doc, text, color="FEF3C7", border="F59E0B"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(8)
    run = p.add_run("ℹ  " + text)
    run.font.size = Pt(9.5)
    run.italic = True
    run.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
    return p

# ════════════════════════════════════════════════════════════════════════
def build_doc():
    doc = Document()

    # Seitenränder
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Deckblatt ─────────────────────────────────────────────────────
    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Plattenwechsler")
    r.font.size = Pt(28); r.bold = True
    r.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("Technische Dokumentation — Raspberry Pi Software")
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)

    doc.add_paragraph()
    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = t3.add_run("DHBW · Projekt G6-TWIE23A · Mai 2026")
    r3.font.size = Pt(11); r3.italic = True
    r3.font.color.rgb = RGBColor(0x6C, 0x72, 0x80)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 1. Systemübersicht
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "1  Systemübersicht")
    add_para(doc,
        "Der Plattenwechsler ist ein automatisierter Roboter, der fertig gedruckte "
        "Druckplatten von bis zu vier 3D-Druckern entnimmt, auf Ablageplätzen zwischenlagert "
        "und neue Druckplatten aus einem Magazin einlegt. Ein Raspberry Pi 4 übernimmt "
        "die gesamte Steuerungslogik, Benutzeroberfläche und Anbindung an externe Dienste. "
        "Ein ESP32 auf dem Schlitten steuert die Motoren und liest die Schlitten-Sensoren aus.")

    add_heading(doc, "1.1  Systemarchitektur", level=2)
    add_para(doc,
        "Die folgende Grafik zeigt die Software-Architektur des Raspberry Pi und seine "
        "Schnittstellen zu externen Komponenten.")
    add_img(doc, make_architektur(), width_cm=16)

    add_para(doc,
        "Der Raspberry Pi kommuniziert mit dem ESP32 über UART (115200 Baud, USB-Serial). "
        "Druckplatten-Fertig-Signale der Drucker werden direkt über GPIO-Eingänge erkannt. "
        "Die Schlitten-Endschalter (X- und Z-Achse) sind ebenfalls am Pi angeschlossen und "
        "werden während der Referenzfahrt per HOME_SWITCH_HIT-Kommando an den ESP weitergeleitet. "
        "MQTT und Telegram ermöglichen Fernzugriff und Statusmeldungen.")

    add_heading(doc, "1.2  Verwendete Technologien", level=2)
    add_table(doc,
        ["Komponente", "Technologie / Bibliothek", "Zweck"],
        [
            ["Betriebssystem",   "Raspberry Pi OS (64-bit)",       "Basis-Betriebssystem"],
            ["Programmiersprache","Python 3.11",                   "Gesamte Pi-Software"],
            ["HMI",              "PyQt5",                          "Touchscreen-Oberfläche"],
            ["Kommunikation ESP","pyserial",                       "UART-Verbindung zum ESP32"],
            ["GPIO",             "RPi.GPIO",                       "Endschalter, Drucker-Pins"],
            ["MQTT",             "paho-mqtt",                      "Fernsteuerung / Status"],
            ["Konfiguration",    "PyYAML",                         "config.yaml lesen/schreiben"],
            ["Telegram",         "python-telegram-bot",            "Bot-Benachrichtigungen"],
        ],
        col_widths=[4.0, 5.0, 6.5]
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 2. Software-Architektur
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "2  Software-Architektur")
    add_para(doc,
        "Die Software ist in klar getrennte Module aufgeteilt. Jedes Modul hat eine "
        "einzige Verantwortung und kommuniziert über definierte Schnittstellen (Callbacks, "
        "Methodenaufrufe) mit den anderen.")

    add_heading(doc, "2.1  Modulübersicht", level=2)
    add_table(doc,
        ["Modul", "Datei", "Aufgabe"],
        [
            ["Hauptablauf",      "core/hauptablauf.py",    "Statemachine, Plattenwechsel-Logik, Service-Modus"],
            ["ESP-Client",       "io_/esp_client.py",      "UART-Kommunikation, Protokoll-Parsing, Kommando-Queue"],
            ["Mock-ESP-Client",  "io_/mock_esp_client.py", "Hardware-Simulation für Tests ohne ESP"],
            ["GPIO-Manager",     "io_/gpio_manager.py",    "Endschalter-Interrupts, Drucker-fertig-Pins"],
            ["MQTT-Client",      "io_/mqtt_client.py",     "Verbindung zu Broker, Publish/Subscribe"],
            ["Telegram-Client",  "io_/telegram_client.py", "Bot-API, Befehle empfangen, Status senden"],
            ["Config",           "config.py",              "YAML laden/speichern, Drucker/Ablage/Magazin-CRUD"],
            ["Types",            "types.py",               "Enums, Dataclasses, Fehlerhierarchie"],
            ["HMI",              "ui/main_window.py",      "PyQt5-Oberfläche, alle Tabs und Dialoge"],
            ["Startpunkt",       "main.py",                "Dependency Injection, Initialisierung, Logging"],
        ],
        col_widths=[3.5, 4.5, 8.5]
    )

    add_heading(doc, "2.2  Threading-Modell", level=2)
    add_para(doc,
        "Die Anwendung nutzt mehrere Threads, um UI-Reaktionsfähigkeit und Hardware-Kommunikation "
        "zu entkoppeln:")
    add_table(doc,
        ["Thread", "Name", "Aufgabe"],
        [
            ["UI-Thread",        "Qt-Main-Thread",   "PyQt5-Ereignisschleife, alle Widget-Updates"],
            ["Worker-Thread",    "Hauptablauf",       "Statemachine, Plattenwechsel-Ausführung (blockierend)"],
            ["Reader-Thread",    "EspReader",         "Kontinuierliches Lesen vom UART, Parsing"],
            ["Watchdog-Thread",  "EspWatchdog",       "Heartbeat-Überwachung, Reconnect bei Verbindungsverlust"],
            ["Heartbeat-Thread", "MockHeartbeat",     "Nur im Mock: periodische HEARTBEAT-Events"],
            ["GPIO-Threads",     "RPi.GPIO intern",   "Interrupt-Callbacks für Endschalter und Drucker-Pins"],
        ],
        col_widths=[3.5, 4.0, 9.0]
    )
    add_info_box(doc,
        "Alle Zugriffe auf gemeinsam genutzte Daten (ESP-Status, Config) sind durch "
        "threading.RLock() abgesichert. UI-Updates aus Hintergrund-Threads erfolgen "
        "ausschließlich über Qt-Signale (thread-safe).")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 3. Statemachine
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "3  Statemachine")
    add_para(doc,
        "Der Hauptablauf des Raspberry Pi wird durch eine Statemachine gesteuert. "
        "Der aktuelle Zustand bestimmt, welche Aktionen erlaubt sind und wie auf "
        "Ereignisse (Druckersignale, Fehler, Benutzereingaben) reagiert wird.")

    add_img(doc, make_statemachine(), width_cm=15.5)

    add_heading(doc, "3.1  Zustandsbeschreibung", level=2)
    add_table(doc,
        ["Zustand", "Bedeutung", "Auslöser"],
        [
            ["INIT",           "System startet, wartet auf ESP-Verbindung",     "Programmstart"],
            ["REFERENZFAHRT",  "Homing aller Achsen läuft",                     "Nach INIT oder nach Fehlerquittierung"],
            ["BEREITSCHAFT",   "System wartet auf Aufträge",                    "Nach erfolgreicher Referenzfahrt"],
            ["PLATTENWECHSEL", "Wechsel-Sequenz läuft aktiv",                   "Auftrag aus Queue entnommen"],
            ["SERVICE",        "Manueller Fahrbetrieb, keine Aufträge",         "Bediener aktiviert Service-Modus"],
            ["FEHLER",         "Fehler aufgetreten, Motoren gestoppt",          "ESP-Fehler, Timeout, interner Fehler"],
            ["NOT_AUS",        "Not-Aus-Taster betätigt, sofortiger Stopp",     "GPIO-Interrupt am Not-Aus-Pin"],
        ],
        col_widths=[3.8, 6.5, 6.2]
    )

    add_heading(doc, "3.2  Fehlerbehandlung", level=2)
    add_para(doc,
        "Tritt während des Betriebs ein Fehler auf, wechselt das System sofort in den "
        "Zustand FEHLER. Der ESP stoppt alle Motoren. Der Bediener muss den Fehler am "
        "Touchscreen quittieren. Danach erfolgt automatisch eine neue Referenzfahrt. "
        "Tritt der Fehler während eines laufenden Plattenwechsels auf, wird die Auftrags-Queue "
        "geleert, da der physische Zustand von Ablage, Magazin und Schlitten unbekannt ist.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 4. Plattenwechsel-Sequenz
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "4  Plattenwechsel-Sequenz")
    add_para(doc,
        "Der Plattenwechsel läuft in fünf Phasen ab. Der Pi sendet Kommandos an den ESP, "
        "wartet jeweils auf die Abschluss-Events und aktualisiert anschließend den Lager-Status "
        "in der config.yaml.")

    add_img(doc, make_sequenz(), width_cm=16)

    add_heading(doc, "4.1  OPEN_DOOR / CLOSE_DOOR", level=2)
    add_para(doc,
        "Das Öffnen und Schließen der Druckertüren erfolgt vollständig durch den ESP. "
        "Der Pi übergibt lediglich die Positionsparameter. Der ESP fährt intern zur "
        "Anfahrposition, führt den Kreisbogen aus und kehrt anschließend zur Ausgangsposition zurück.")
    add_para(doc, "Für CLOSE_DOOR berechnet der Pi die Anfahrposition X selbst:", bold=True)
    add_mono(doc, "x_close = pos_x + int(tuer_radius × (cos(tuer_winkel°) − 1))")
    add_para(doc,
        "Dieses Ergebnis entspricht der X-Position am Ende des Öffnungsbogens "
        "und wird als x_approach an den ESP übermittelt.")

    add_heading(doc, "4.2  Lager-Status-Update", level=2)
    add_para(doc,
        "Nach jeder Plattenbewegung wird der Status in der config.yaml aktualisiert:")
    add_table(doc,
        ["Ereignis", "Aktion", "Auswirkung"],
        [
            ["DEPOSIT in Ablage abgeschlossen",  "ablage_belegt_setzen(id, True)",        "Ablage als BELEGT markiert"],
            ["PICKUP aus Magazin abgeschlossen",  "magazin_verfuegbar_setzen(id, False)",  "Magazin als LEER markiert"],
            ["Bediener quittiert manuell",        "Toggle-Button im Lager-Tab",            "Status direkt umschalten"],
        ],
        col_widths=[5.5, 5.5, 5.5]
    )
    add_info_box(doc,
        "Vor dem Plattenwechsel prüft der Pi, ob eine freie Ablage und ein verfügbares "
        "Magazin vorhanden sind. Fehlt eines davon, wird der Auftrag sofort mit "
        "ABLAGE_VOLL bzw. MAGAZIN_LEER abgelehnt — kein Motor startet.")

    add_img(doc, make_lager(), width_cm=14)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 5. Schnittstelle Pi ↔ ESP32
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "5  Schnittstelle Pi ↔ ESP32")
    add_para(doc,
        "Die Kommunikation zwischen Pi und ESP32 erfolgt über UART (USB-Serial, 115200 Baud, "
        "ASCII, eine Nachricht pro Zeile, Felder getrennt durch Semikolon).")

    add_heading(doc, "5.1  Kommandos (Pi → ESP)", level=2)
    add_para(doc, "Jedes Kommando folgt dem Schema:", bold=True)
    add_mono(doc, "CMD;<id>;<befehl>[;<key>=<value>...]")
    add_para(doc,
        "Die ID ist eine vom Pi vergebene Ganzzahl ≥ 0. Der ESP spiegelt sie in allen "
        "Antworten zurück, sodass Antworten eindeutig zugeordnet werden können.")

    add_table(doc,
        ["Kommando", "Parameter", "Voraussetzung"],
        [
            ["PING",           "—",                                           "immer"],
            ["STATUS",         "—",                                           "immer"],
            ["STREAM_ON/OFF",  "—",                                           "immer"],
            ["STOP",           "—",                                           "immer"],
            ["HOME",           "—",                                           "nicht ERROR, nicht busy"],
            ["MOVE_HOME",      "—",                                           "READY oder STOPPED, referenziert"],
            ["MOVE_TO",        "x=<mm>, z=<mm>",                             "READY oder STOPPED, referenziert"],
            ["HOME_SWITCH_HIT","axis=X|Z",                                   "nur in BUSY_HOMING / BUSY_MOVE_HOME"],
            ["OPEN_DOOR",      "x_approach, z_approach, arm_extend, radius, angle, hook_drop", "READY, referenziert"],
            ["CLOSE_DOOR",     "x_approach, z_approach, arm_extend, radius, angle, hook_drop", "READY, referenziert"],
            ["PICKUP",         "gripper_depth=<mm>, lift_offset=<mm>",       "READY, referenziert"],
            ["DEPOSIT",        "gripper_depth=<mm>, lift_offset=<mm>",       "READY, referenziert"],
            ["RESET_ERROR",    "—",                                           "nur in ERROR"],
        ],
        col_widths=[3.2, 6.8, 5.5]
    )

    add_heading(doc, "5.2  Antworten (ESP → Pi)", level=2)
    add_para(doc,
        "Der ESP antwortet auf jedes Kommando sofort mit RSP (Sofortantwort) und "
        "sendet nach Abschluss einer Aktion ein EVT (Event):")
    add_mono(doc, "RSP;<id>;ACK              ← Kommando akzeptiert")
    add_mono(doc, "RSP;<id>;ERR;<fehlercode> ← Kommando abgelehnt")
    add_mono(doc, "EVT;<id>;OK;<event_name>;x=<mm>;z=<mm>  ← Aktion abgeschlossen")
    add_mono(doc, "EVT;0;STATE;<zustand>;ref=<0|1>;x=<mm>;z=<mm>  ← Zustandswechsel")
    add_mono(doc, "EVT;0;STATUS;state=…;error=…;ref=…;x=…;z=…;…  ← vollständiger Status")
    add_mono(doc, "EVT;0;HEARTBEAT;uptime_ms=<ms>;state=…  ← alle 1000 ms")
    add_mono(doc, "EVT;0;ERR;<fehlercode>;x=<mm>;z=<mm>    ← Fehler (→ ERROR-Zustand)")

    add_heading(doc, "5.3  Besonderheit: Endschalter", level=2)
    add_para(doc,
        "Die Schlitten-Endschalter für X- und Z-Achse sind am Raspberry Pi angeschlossen. "
        "Der Pi erkennt den Auslöse-Impuls per GPIO-Interrupt und leitet ihn sofort als "
        "HOME_SWITCH_HIT-Kommando an den ESP weiter. Greifer- und Türarm-Endschalter "
        "sitzen auf dem Schlitten und werden direkt vom ESP ausgewertet.")

    add_heading(doc, "5.4  Kommunikationsablauf", level=2)
    add_img(doc, make_protokoll(), width_cm=15.5)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 6. Konfiguration
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "6  Konfiguration")
    add_para(doc,
        "Alle Einstellungen werden in der Datei config.yaml gespeichert. "
        "Änderungen über die HMI werden sofort in die Datei zurückgeschrieben "
        "und sind damit dauerhaft gespeichert.")

    add_heading(doc, "6.1  ESP-Parameter", level=2)
    add_table(doc,
        ["Parameter", "Standardwert", "Bedeutung"],
        [
            ["port",               "/dev/ttyUSB0", "Serieller Port zum ESP32"],
            ["baud",               "115200",        "Baudrate"],
            ["ack_timeout_s",      "1.0",           "Maximale Wartezeit auf RSP-Antwort"],
            ["move_timeout_s",     "60.0",          "Timeout für Fahrbewegungen und OPEN/CLOSE_DOOR"],
            ["home_timeout_s",     "30.0",          "Timeout für Referenzfahrt und MOVE_HOME"],
            ["mech_timeout_s",     "10.0",          "Timeout für PICKUP und DEPOSIT"],
            ["heartbeat_timeout_s","5.0",           "Nach dieser Zeit ohne HEARTBEAT: Reconnect"],
            ["skip_homing",        "false",         "Referenzfahrt überspringen (nur für Tests)"],
        ],
        col_widths=[4.5, 3.5, 8.5]
    )

    add_heading(doc, "6.2  Drucker-Konfiguration", level=2)
    add_para(doc,
        "Jeder Drucker besitzt einen eigenen Eintrag unter dem Schlüssel drucker. "
        "Alle Werte sind über den Drucker-Tab der HMI änderbar:")
    add_table(doc,
        ["Feld", "Bedeutung"],
        [
            ["id",              "Eindeutige Drucker-ID (1–n)"],
            ["name",            "Anzeigename in der HMI"],
            ["pin_fertig",      "GPIO-BCM-Pin: LOW = Drucker fertig. 0 = deaktiviert"],
            ["pos_x",           "X-Position des Druckers in mm (Ausgangsposition + x_approach für OPEN_DOOR)"],
            ["pos_z_anfahr",    "Z-Höhe der Ausgangsposition vor dem Drucker in mm"],
            ["pos_z_tuer",      "Z-Höhe für OPEN/CLOSE_DOOR (z_approach) in mm"],
            ["pos_z_druckbett", "Z-Höhe der Gabel-Bereitschaftsposition in mm"],
            ["door_arm_hub_mm", "Wie weit der Türarm ausfährt (arm_extend) in mm"],
            ["tuer_radius",     "Radius des Türscharniers für den Kreisbogen in mm"],
            ["tuer_winkel",     "Öffnungswinkel der Tür in Grad"],
            ["gripper_depth",   "Greifer-Ausfahrtiefe für PICKUP/DEPOSIT in mm"],
            ["lift_offset",     "Z-Hub beim PICKUP/DEPOSIT in mm"],
        ],
        col_widths=[4.0, 12.5]
    )

    add_heading(doc, "6.3  Ablagen und Magazine", level=2)
    add_para(doc,
        "Ablagen und Magazine werden als Listen konfiguriert. "
        "Jeder Eintrag hat eine Position (x, z), Greifer-Parameter und einen Status:")
    add_table(doc,
        ["Feld", "Ablage", "Magazin"],
        [
            ["id, name",       "Eindeutig, Anzeigename",              "Eindeutig, Anzeigename"],
            ["x, z",           "Position in mm",                      "Position in mm"],
            ["gripper_depth",  "Greifer-Ausfahrtiefe in mm",          "Greifer-Ausfahrtiefe in mm"],
            ["lift_offset",    "Z-Hub in mm",                         "Z-Hub in mm"],
            ["belegt",         "true = Platte liegt auf Ablage",      "—"],
            ["verfuegbar",     "—",                                   "true = Platte vorhanden"],
        ],
        col_widths=[3.5, 7.0, 7.0]
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 7. HMI-Bedienung
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "7  HMI-Bedienung")
    add_para(doc,
        "Die Benutzeroberfläche läuft im Vollbild auf einem 7-Zoll-Touchscreen (1280×720). "
        "Die Navigationsleiste links enthält fünf Tabs. Eine QWERTZ-Bildschirmtastatur "
        "erscheint automatisch bei Texteingabe-Feldern.")

    add_heading(doc, "7.1  Tab-Übersicht", level=2)
    add_table(doc,
        ["Tab", "Inhalt", "Aktionen"],
        [
            ["Status",    "Systemzustand, ESP-Status, aktive Fehler, Queue-Länge",
                          "Fehler quittieren, Auftrag manuell starten"],
            ["Manuell",   "Drucker-Kacheln mit aktuellem Status",
                          "Auftrag manuell für einzelnen Drucker aufgeben"],
            ["Service",   "Schlitten-Fahrbefehle, Referenzfahrt",
                          "Service-Modus aktivieren, zu Drucker/Ablage/Magazin fahren"],
            ["Drucker",   "Liste aller konfigurierten Drucker",
                          "Drucker hinzufügen, bearbeiten, löschen"],
            ["Lager",     "Ablagen und Magazine mit Status",
                          "Status umschalten (frei/belegt, verfügbar/leer), Position bearbeiten"],
        ],
        col_widths=[2.5, 6.5, 7.5]
    )

    add_heading(doc, "7.2  Drucker-Konfiguration im HMI", level=2)
    add_para(doc,
        "Der Drucker-Editor öffnet sich als Vollbild-Dialog. Er ist in zwei Tabs unterteilt:")
    add_para(doc, "• Positionen: X-Position, Z Anfahrt, Z Druckbett", italic=True)
    add_para(doc, "• Tür / Greifer: Z Türarm, Türarm-Hub, Türradius, Greifer, Hub-Offset, Öffnungswinkel", italic=True)
    add_para(doc,
        "Name und Pin Fertig befinden sich immer sichtbar über den Tabs. "
        "Jeder Pin Fertig darf nur einmal vergeben werden — bei Duplikat erscheint eine Warnung.")

    add_heading(doc, "7.3  Service-Modus", level=2)
    add_para(doc,
        "Im Service-Modus können einzelne Positionen manuell angefahren werden. "
        "Eingehende Drucker-fertig-Signale werden gespeichert und nach Verlassen des "
        "Service-Modus automatisch in die Queue eingereiht. "
        "Der Service-Modus ist nur aus dem Zustand BEREITSCHAFT erreichbar.")

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 8. Fehlerbehandlung
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "8  Fehlerbehandlung")
    add_para(doc,
        "Alle Fehler werden durch die Klasse FehlerBehandlung verwaltet. "
        "Es kann immer nur ein aktiver Fehler vorliegen. Fehler müssen am Touchscreen "
        "quittiert werden, bevor der Betrieb fortgesetzt wird.")

    add_heading(doc, "8.1  Fehlerklassen", level=2)
    add_table(doc,
        ["Fehlerklasse", "Ursache", "Typische ESP-Codes"],
        [
            ["Kommunikationsfehler", "UART-Verbindung unterbrochen, Timeout auf Antwort",   "INVALID_COMMAND, BUSY, INVALID_STATE"],
            ["Fahrfehler",           "Zielposition nicht erreicht, Timeout, kein Referenz", "MOVE_TIMEOUT, HOMING_TIMEOUT, NOT_REFERENCED"],
            ["Hindernis",            "TF-Luna LiDAR erkennt Objekt im Weg",                "OBSTACLE"],
            ["Sensorfehler",         "Sensor ausgefallen oder nicht initialisierbar",       "SENSOR_FAULT_OBSTACLE, SENSOR_FAULT_GRIPPER"],
            ["Tuerfehler",           "Türsensor: Tür nicht offen, Türarm-Problem",         "DOOR_NOT_OPEN"],
            ["Entnahmefehler",       "Platte nach PICKUP nicht erkannt",                   "PLATE_NOT_DETECTED"],
            ["Interner Fehler",      "Unerwartete Ausnahme in der Python-Software",        "—"],
            ["Ablage voll",          "Keine freie Ablage vor Plattenwechsel",              "—"],
            ["Magazin leer",         "Kein verfügbares Magazin vor Plattenwechsel",        "—"],
            ["Not-Aus",              "Not-Aus-Taster betätigt",                            "—"],
        ],
        col_widths=[4.0, 6.5, 6.0]
    )

    add_heading(doc, "8.2  Verhalten im Fehlerfall", level=2)
    add_table(doc,
        ["Schritt", "Aktion"],
        [
            ["1", "ESP stoppt sofort alle Motoren"],
            ["2", "Pi wechselt in Zustand FEHLER"],
            ["3", "Fehler wird in HMI, Telegram und MQTT gemeldet"],
            ["4", "Bediener quittiert Fehler am Touchscreen"],
            ["5", "Pi sendet RESET_ERROR an ESP"],
            ["6", "Pi führt neue Referenzfahrt durch (→ REFERENZFAHRT)"],
            ["7", "Normalbetrieb wird aufgenommen (→ BEREITSCHAFT)"],
        ],
        col_widths=[1.5, 15.0]
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════
    # 9. MQTT & Telegram
    # ══════════════════════════════════════════════════════════════════
    add_heading(doc, "9  Fernzugriff: MQTT und Telegram")

    add_heading(doc, "9.1  MQTT", level=2)
    add_para(doc,
        "Der MQTT-Client verbindet sich zu einem konfigurierbaren Broker (Standard: HiveMQ). "
        "Der Basis-Topic ist in config.yaml festgelegt. Status-Nachrichten werden "
        "mit retain=true veröffentlicht, sodass neue Subscriber sofort den letzten Status erhalten.")
    add_table(doc,
        ["Topic (relativ zu base_topic)", "Richtung", "Inhalt"],
        [
            ["status",        "Pi → Broker",  "Systemzustand, ESP-Status, Queue, Fehler (JSON)"],
            ["auftrag",       "Broker → Pi",  "Drucker-ID als Integer: Auftrag aufgeben"],
            ["stop",          "Broker → Pi",  "Manuellen Stopp auslösen"],
        ],
        col_widths=[5.5, 3.0, 8.0]
    )

    add_heading(doc, "9.2  Telegram", level=2)
    add_para(doc,
        "Der Telegram-Bot erlaubt autorisierten Nutzern (konfigurierte Chat-IDs) "
        "Statusabfragen und Steuerungsbefehle per Nachricht:")
    add_table(doc,
        ["Befehl", "Funktion"],
        [
            ["/status",    "Aktuellen Systemzustand und Fehler abfragen"],
            ["/auftrag N", "Auftrag für Drucker N aufgeben"],
            ["/stop",      "Manuellen Stopp auslösen"],
        ],
        col_widths=[3.5, 13.0]
    )
    add_para(doc,
        "Status-Nachrichten werden automatisch nach jedem Plattenwechsel und bei Fehlern "
        "an den ersten konfigurierten Chat gesendet (send_status_to_first: true).")

    # Abschluss
    doc.add_page_break()
    add_heading(doc, "Anhang — Hinweise")
    add_para(doc,
        "Die vollständige Schnittstellenbeschreibung (Protokoll-Details, alle Fehlercodes, "
        "ESP-Zustandscodes, GPIO-Pinbelegung Pi und ESP32, config.yaml-Vorlage sowie "
        "die Installationsanleitung) befindet sich im separaten Anhang-Dokument.")
    add_para(doc,
        "Den ESP32-Teil der Dokumentation (Firmware-Architektur, Motoransteuerung, "
        "Sensor-Integration) übernimmt der zuständige Kollege.", italic=True)

    out = "/home/nura/plattenwechsler/docs/Plattenwechsler_Doku_Pi.docx"
    doc.save(out)
    print(f"Gespeichert: {out}")

if __name__ == "__main__":
    build_doc()
