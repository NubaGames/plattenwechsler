"""Erstellt Schnittstellen_Doku als Word-Datei."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    col = {1: RGBColor(0xC8, 0x10, 0x2E),
           2: RGBColor(0x1E, 0x20, 0x28),
           3: RGBColor(0x2D, 0x31, 0x42)}.get(level, RGBColor(0x1E, 0x20, 0x28))
    for run in h.runs:
        run.font.color.rgb = col
    h.paragraph_format.space_before = Pt({1: 16, 2: 10, 3: 7}.get(level, 6))
    h.paragraph_format.space_after  = Pt(5)
    return h


def add_para(doc, text, bold=False, italic=False, size=10.5, indent=0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold; run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.space_after = Pt(2)
    if level:
        p.paragraph_format.left_indent = Cm(level * 0.5)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_mono(doc, text, indent=0.8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(1)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(8)
    run = p.add_run("▶  " + text)
    run.font.size = Pt(9.5)
    run.italic = True
    run.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        cell.text = h
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "C8102E"); shd.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shd)
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        trow = table.rows[ri + 1]
        for ci, val in enumerate(row):
            cell = trow.cells[ci]; cell.text = str(val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                if ci == 0:
                    run.font.name = "Courier New"
        if ri % 2 == 0:
            for ci in range(len(headers)):
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F8F9FA"); shd.set(qn("w:val"), "clear")
                trow.cells[ci]._tc.get_or_add_tcPr().append(shd)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C8102E")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ════════════════════════════════════════════════════════════════════════
def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Deckblatt ────────────────────────────────────────────────────
    doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Schnittstellendokumentation")
    r.font.size = Pt(26); r.bold = True
    r.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)

    t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("Plattenwechsler — G6-TWIE23A")
    r2.font.size = Pt(15); r2.font.color.rgb = RGBColor(0x1E, 0x20, 0x28)

    doc.add_paragraph()
    t3 = doc.add_paragraph(); t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = t3.add_run("DHBW Stuttgart · IoT 2026 · Stand: Mai 2026")
    r3.font.size = Pt(11); r3.italic = True
    r3.font.color.rgb = RGBColor(0x6C, 0x72, 0x80)

    doc.add_paragraph()
    toc = doc.add_paragraph(); toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = toc.add_run(
        "Inhalt: ESP32 ↔ Pi UART-Protokoll  ·  MQTT  ·  Web-App API  ·  GPIO  ·  config.yaml")
    r4.font.size = Pt(10); r4.font.color.rgb = RGBColor(0x6C, 0x72, 0x80)
    r4.italic = True

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # 1. ESP32 ↔ Pi: UART-Protokoll
    # ════════════════════════════════════════════════════════════════
    add_heading(doc, "1  ESP32 ↔ Pi: UART-Protokoll")

    add_heading(doc, "1.1  Verbindung", level=2)
    add_table(doc,
        ["Parameter", "Wert"],
        [
            ["Schnittstelle",    "UART over USB (/dev/ttyUSB0)"],
            ["Baudrate",         "115200, 8N1"],
            ["Encoding",         "ASCII, eine Nachricht pro Zeile (\\n oder \\r\\n)"],
            ["Feldtrenner",      "Semikolon  ;"],
            ["Positionsangaben", "Immer in Millimeter (Ganzzahl)"],
            ["Max. Zeilenlänge", "160 Zeichen"],
            ["Kommando-Queue",   "8 Einträge; bei Überlauf: RSP;<id>;ERR;BUSY"],
        ],
        col_widths=[5.0, 11.5]
    )

    add_divider(doc)

    # 1.2 Kommandos
    add_heading(doc, "1.2  Pi → ESP: Kommandos", level=2)
    add_para(doc, "Alle Kommandos folgen dem Schema:")
    add_mono(doc, "CMD;<id>;<befehl>[;<key>=<value>...]")
    add_para(doc,
        "<id> ist eine vom Pi vergebene Ganzzahl ≥ 0. Der ESP spiegelt sie in allen "
        "Antworten zurück.")
    doc.add_paragraph()

    add_table(doc,
        ["Kommando", "Syntax", "Voraussetzung"],
        [
            ["PING",            "CMD;<id>;PING",                                         "immer"],
            ["STATUS",          "CMD;<id>;STATUS",                                       "immer"],
            ["STREAM_ON",       "CMD;<id>;STREAM_ON",                                    "immer"],
            ["STREAM_OFF",      "CMD;<id>;STREAM_OFF",                                   "immer"],
            ["STOP",            "CMD;<id>;STOP",                                         "immer"],
            ["HOME",            "CMD;<id>;HOME",                                         "nicht ERROR, nicht busy"],
            ["MOVE_HOME",       "CMD;<id>;MOVE_HOME",                                    "READY oder STOPPED, referenziert"],
            ["MOVE_TO",         "CMD;<id>;MOVE_TO;x=<mm>;z=<mm>",                       "READY oder STOPPED, referenziert"],
            ["RESET_ERROR",     "CMD;<id>;RESET_ERROR",                                  "nur in ERROR"],
            ["HOME_SWITCH_HIT", "CMD;<id>;HOME_SWITCH_HIT;axis=<X|Z>",                  "nur in BUSY_HOMING / BUSY_MOVE_HOME"],
            ["OPEN_DOOR",       "CMD;<id>;OPEN_DOOR;x_approach=<mm>;z_approach=<mm>;\n  arm_extend=<mm>;radius=<mm>;angle=<deg>;hook_drop=<mm>",
             "READY oder STOPPED, referenziert"],
            ["CLOSE_DOOR",      "CMD;<id>;CLOSE_DOOR;x_approach=<mm>;z_approach=<mm>;\n  arm_extend=<mm>;radius=<mm>;angle=<deg>;hook_drop=<mm>",
             "READY oder STOPPED, referenziert"],
            ["PICKUP",          "CMD;<id>;PICKUP;gripper_depth=<mm>;lift_offset=<mm>",   "READY oder STOPPED, referenziert"],
            ["DEPOSIT",         "CMD;<id>;DEPOSIT;gripper_depth=<mm>;lift_offset=<mm>",  "READY oder STOPPED, referenziert"],
        ],
        col_widths=[3.5, 7.5, 5.5]
    )

    add_divider(doc)

    # OPEN_DOOR
    add_heading(doc, "OPEN_DOOR — Parameter", level=3)
    add_para(doc,
        "Der Pi berechnet x_approach selbst (= pos_x des Druckers). "
        "Der ESP führt die komplette Bogenbewegung intern aus.")
    add_para(doc, "Geometrie:  Δx(θ) = radius · (cos(θ) − 1)", bold=True)
    add_para(doc, "ESP-interner Ablauf:", bold=True)
    add_numbered(doc, "Schlitten fährt auf (x_approach, z_approach)")
    add_numbered(doc, "Türarm fährt arm_extend mm aus")
    add_numbered(doc, "Z fährt hook_drop mm nach oben (einhaken)")
    add_numbered(doc, "Kreisbogen öffnen: Arm konstante Geschwindigkeit, X-Achse folgt")
    add_numbered(doc, "Z fährt hook_drop mm nach unten (aushaken)")
    add_numbered(doc, "Türarm fährt ein")
    add_numbered(doc, "Schlitten kehrt zur Ausgangsposition zurück")
    doc.add_paragraph()

    add_table(doc,
        ["Parameter", "Typ", "Bedeutung"],
        [
            ["x_approach", "int (mm)", "Absolute X-Anfahrposition"],
            ["z_approach", "int (mm)", "Absolute Z-Anfahrposition"],
            ["arm_extend",  "int (mm)", "Ausfahrlänge Türarm"],
            ["radius",      "int (mm)", "Kreisradius (Scharnier → Greifpunkt)"],
            ["angle",       "int (°)",  "Öffnungswinkel"],
            ["hook_drop",   "int (mm)", "Z-Versatz Einhakmechanismus; 0 = kein Versatz"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )
    add_mono(doc, "CMD;7;OPEN_DOOR;x_approach=370;z_approach=105;arm_extend=30;radius=150;angle=160;hook_drop=15")
    doc.add_paragraph()

    add_divider(doc)

    # CLOSE_DOOR
    add_heading(doc, "CLOSE_DOOR — Parameter", level=3)
    add_para(doc,
        "Der Schlitten steht in der Ausgangsposition (Tür ist offen). Der Pi berechnet "
        "die Anfahrposition x_approach aus der OPEN_DOOR-Geometrie:")
    add_mono(doc, "x_close_approach = x_open_approach + radius · (cos(angle) − 1)")
    add_para(doc,
        "Für Winkel > 0° ist cos(angle) < 1, d.h. x_close < x_open "
        "(Schlitten näher am Drucker).")
    add_para(doc, "ESP-interner Ablauf:", bold=True)
    add_numbered(doc, "Schlitten fährt auf (x_approach, z_approach)")
    add_numbered(doc, "Türarm fährt auf Grifftiefe bei geöffneter Tür: arm_extend + radius · sin(angle) mm")
    add_numbered(doc, "Z fährt hook_drop mm nach oben (einhaken)")
    add_numbered(doc, "Kreisbogen schließen (rückwärts)")
    add_numbered(doc, "Z fährt hook_drop mm nach unten (aushaken)")
    add_numbered(doc, "Türarm fährt ein")
    add_numbered(doc, "Schlitten kehrt zur Ausgangsposition zurück")
    doc.add_paragraph()

    add_para(doc, "Parameter: identisch mit OPEN_DOOR (gleiche Werte übergeben).")
    add_mono(doc, "CMD;10;CLOSE_DOOR;x_approach=79;z_approach=105;arm_extend=30;radius=150;angle=160;hook_drop=15")
    add_para(doc,
        "Beispiel: angle=160°, radius=150, x_open=370 → x_close ≈ 79 mm",
        italic=True, size=9)
    doc.add_paragraph()

    add_divider(doc)

    # PICKUP
    add_heading(doc, "PICKUP — Parameter", level=3)
    add_para(doc,
        "Der Schlitten steht bereits auf der Zielposition (vor Drucker oder Stellplatz).")
    add_para(doc, "ESP-interner Ablauf:", bold=True)
    add_numbered(doc, "Greifer fährt gripper_depth mm aus")
    add_numbered(doc, "Z-Achse hebt um lift_offset mm → Platte auf der Gabel")
    add_numbered(doc, "Greifer fährt ein")
    doc.add_paragraph()
    add_table(doc,
        ["Parameter", "Typ", "Bedeutung"],
        [
            ["gripper_depth", "int (mm)", "Ausfahrtiefe des Greifers"],
            ["lift_offset",   "int (mm)", "Anhebung nach dem Ausfahren"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )
    add_mono(doc, "CMD;5;PICKUP;gripper_depth=120;lift_offset=8")
    doc.add_paragraph()

    add_divider(doc)

    # DEPOSIT
    add_heading(doc, "DEPOSIT — Parameter", level=3)
    add_para(doc,
        "Der Schlitten trägt eine Platte und steht auf der Zielposition.")
    add_para(doc, "ESP-interner Ablauf:", bold=True)
    add_numbered(doc, "Z-Achse hebt um lift_offset mm → Platte über Stellfläche")
    add_numbered(doc, "Greifer fährt gripper_depth mm aus")
    add_numbered(doc, "Z-Achse senkt um lift_offset mm → Platte liegt auf")
    add_numbered(doc, "Greifer fährt ein")
    doc.add_paragraph()
    add_table(doc,
        ["Parameter", "Typ", "Bedeutung"],
        [
            ["gripper_depth", "int (mm)", "Ausfahrtiefe des Greifers"],
            ["lift_offset",   "int (mm)", "Anhebung vor dem Ausfahren"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )
    add_mono(doc, "CMD;6;DEPOSIT;gripper_depth=120;lift_offset=8")
    doc.add_paragraph()

    add_divider(doc)

    # HOME_SWITCH_HIT
    add_heading(doc, "HOME_SWITCH_HIT", level=3)
    add_para(doc,
        "Der Pi meldet einen ausgelösten Endschalter. Nur gültig in "
        "BUSY_HOMING und BUSY_MOVE_HOME.")
    add_note(doc,
        "Die Schlitten-Endschalter (X, Z) sind am Pi angeschlossen. "
        "Greifer- und Türarm-Endschalter sitzen am ESP und werden intern ausgewertet.")
    add_mono(doc, "CMD;2;HOME_SWITCH_HIT;axis=X")
    add_para(doc, "Homing-Reihenfolge:", bold=True)
    add_numbered(doc, "X-Achse fährt auf Endschalter → HOME_SWITCH_HIT;axis=X vom Pi")
    add_numbered(doc, "Z-Achse fährt auf Endschalter → HOME_SWITCH_HIT;axis=Z vom Pi")
    add_numbered(doc, "Greifer + Türarm referenzieren intern parallel")
    add_numbered(doc, "HOME_DONE erst wenn alle vier Achsen referenziert sind")
    doc.add_paragraph()

    doc.add_page_break()

    # 1.3 Antworten
    add_heading(doc, "1.3  ESP → Pi: Antworten", level=2)

    add_heading(doc, "RSP — Sofortantwort", level=3)
    add_para(doc,
        "Kommt direkt nach Empfang des Kommandos, bevor die Aktion abgeschlossen ist.")
    add_mono(doc, "RSP;<id>;ACK               ← Kommando akzeptiert, Aktion läuft")
    add_mono(doc, "RSP;<id>;ERR;<fehlercode>  ← Kommando abgelehnt, kein Motor gestartet")
    doc.add_paragraph()

    add_heading(doc, "EVT — Ereignisse", level=3)
    add_mono(doc, "EVT;<id>;OK;<event_name>;<felder>      ← Aktion abgeschlossen")
    add_mono(doc, "EVT;0;STATE;<zustand>;ref=<0|1>;x=<mm>;z=<mm>")
    add_mono(doc, "EVT;0;STATUS;<alle Statusfelder>")
    add_mono(doc, "EVT;0;ERR;<fehlercode>;x=<mm>;z=<mm>")
    add_mono(doc, "EVT;0;HEARTBEAT;uptime_ms=<ms>;state=<zustand>;x=<mm>;z=<mm>")
    doc.add_paragraph()

    add_table(doc,
        ["event_name", "Auslöser"],
        [
            ["PONG",            "Antwort auf PING"],
            ["HOME_DONE",       "Referenzfahrt abgeschlossen"],
            ["MOVE_HOME_DONE",  "Heimfahrt abgeschlossen"],
            ["MOVE_DONE",       "Zielposition erreicht"],
            ["STOPPED",         "STOP ausgeführt"],
            ["ERROR_RESET",     "RESET_ERROR ausgeführt"],
            ["STREAM_ON",       "Stream eingeschaltet"],
            ["STREAM_OFF",      "Stream ausgeschaltet"],
            ["DOOR_OPEN_DONE",  "Tür geöffnet, Schlitten auf Ausgangsposition"],
            ["DOOR_CLOSE_DONE", "Tür geschlossen, Türarm eingefahren"],
            ["PICKUP_DONE",     "Plattenentnahme abgeschlossen"],
            ["DEPOSIT_DONE",    "Plattenablage abgeschlossen"],
        ],
        col_widths=[4.5, 12.0]
    )

    add_heading(doc, "EVT STATUS — Vollständiger Snapshot", level=3)
    add_para(doc,
        "Gesendet: auf STATUS-Kommando, bei jedem ERROR-Eintritt, "
        "periodisch wenn Stream aktiv (alle 100 ms).")
    add_table(doc,
        ["Feld", "Typ", "Bedeutung"],
        [
            ["state",         "enum",     "Aktueller Zustand"],
            ["error",         "string",   "Aktiver Fehlercode (NONE wenn kein Fehler)"],
            ["ref",           "0/1",      "1 = referenziert"],
            ["x, z",          "int (mm)", "Ist-Position"],
            ["target_x, target_z", "int (mm)", "Soll-Position"],
            ["busy",          "0/1",      "Bewegung aktiv"],
            ["gripper_home",  "0/1",      "Greifer in Heimposition"],
            ["door_arm_home", "0/1",      "Türarm in Heimposition"],
            ["obstacle_ok",   "0/1",      "TF-Luna gesund und frei"],
            ["door_open",     "0/1",      "Druckertür offen (VL53L0X-Auswertung)"],
            ["door_dist_mm",  "int (mm)", "Rohwert Türsensor (Debugging)"],
            ["plate_detected","0/1",      "Plattenerkennungs-Taster aktiv"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )

    add_divider(doc)

    # 1.4 Zustandscodes
    add_heading(doc, "1.4  Zustandscodes", level=2)
    add_table(doc,
        ["Code", "Bedeutung"],
        [
            ["NOT_REFERENCED",  "Bereit, Referenzfahrt noch nicht durchgeführt"],
            ["READY",           "Referenziert, wartet auf Kommando"],
            ["BUSY_HOMING",     "Referenzfahrt läuft"],
            ["BUSY_SCANNING",   "Z-Scan vor Fahrt aus Home-Position"],
            ["BUSY_MOVING",     "Fahrt zu Zielposition"],
            ["BUSY_MOVE_HOME",  "Heimfahrt läuft"],
            ["BUSY_PICKUP",     "Plattenentnahme läuft"],
            ["BUSY_DEPOSIT",    "Plattenablage läuft"],
            ["BUSY_OPEN_DOOR",  "Türöffnung läuft"],
            ["BUSY_CLOSE_DOOR", "Türschließung läuft"],
            ["STOPPED",         "Per STOP angehalten"],
            ["ERROR",           "Fehler, alle Motoren gestoppt, wartet auf RESET_ERROR"],
        ],
        col_widths=[4.5, 12.0]
    )
    add_para(doc, "Übergänge:", bold=True)
    add_bullet(doc, "STOP → STOPPED aus jedem Zustand außer ERROR")
    add_bullet(doc, "STOPPED verhält sich wie READY (HOME, MOVE_TO, MOVE_HOME, PICKUP, DEPOSIT, OPEN/CLOSE_DOOR erlaubt)")
    add_bullet(doc, "Jeder Fehler aus einem BUSY_*-Zustand → sofort ERROR")
    add_bullet(doc, "RESET_ERROR → NOT_REFERENCED, Referenzierung gelöscht")
    doc.add_paragraph()

    add_divider(doc)

    # 1.5 Fehlercodes
    add_heading(doc, "1.5  Fehlercodes", level=2)
    add_table(doc,
        ["Code", "Klasse", "Bedeutung"],
        [
            ["INVALID_COMMAND",       "Kommunikation", "Unbekanntes Kommando oder Syntaxfehler"],
            ["BUSY",                  "Kommunikation", "Kommando-Queue voll"],
            ["INVALID_STATE",         "Kommunikation", "Kommando im aktuellen Zustand nicht erlaubt"],
            ["NOT_REFERENCED",        "Fahrt",         "MOVE_TO ohne vorherige Referenzfahrt"],
            ["MOVE_TIMEOUT",          "Fahrt",         "Zielposition nicht rechtzeitig erreicht"],
            ["HOMING_TIMEOUT",        "Fahrt",         "Referenzfahrt-Timeout abgelaufen"],
            ["POSITION_ERROR",        "Fahrt",         "Rücklese-Position außerhalb Toleranz"],
            ["OBSTACLE",              "Hindernis",     "TF-Luna unterschreitet Stoppabstand"],
            ["SENSOR_FAULT_OBSTACLE", "Sensor",        "TF-Luna defekt oder nicht initialisierbar"],
            ["SENSOR_FAULT_GRIPPER",  "Sensor",        "Greifer-Endschalter antwortet nicht erwartet"],
            ["DRIVER_FAULT",          "Motor",         "Closed-Loop-Treiber (CL42T) meldet Alarm"],
            ["PLATE_NOT_DETECTED",    "Greifer",       "Plattenerkennungs-Taster nach PICKUP nicht ausgelöst"],
            ["DOOR_NOT_OPEN",         "Tür",           "Türsensor: zu geringe Distanz vor PICKUP/DEPOSIT"],
        ],
        col_widths=[4.2, 3.0, 9.3]
    )

    add_divider(doc)

    # 1.6 Kommunikationssequenzen
    add_heading(doc, "1.6  Typische Kommunikationssequenzen", level=2)

    add_heading(doc, "Referenzfahrt", level=3)
    for line in [
        "→ CMD;1;HOME",
        "← RSP;1;ACK",
        "← EVT;0;STATE;BUSY_HOMING;ref=0;x=0;z=0",
        "  [X-Achse fährt auf Endschalter]",
        "→ CMD;2;HOME_SWITCH_HIT;axis=X",
        "← RSP;2;ACK",
        "  [Z-Achse fährt auf Endschalter]",
        "→ CMD;3;HOME_SWITCH_HIT;axis=Z",
        "← RSP;3;ACK",
        "← EVT;1;OK;HOME_DONE;x=0;z=0",
        "← EVT;0;STATE;READY;ref=1;x=0;z=0",
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    add_heading(doc, "Tür öffnen (OPEN_DOOR)", level=3)
    for line in [
        "→ CMD;7;OPEN_DOOR;x_approach=370;z_approach=105;arm_extend=30;radius=150;angle=160;hook_drop=15",
        "← RSP;7;ACK",
        "← EVT;0;STATE;BUSY_OPEN_DOOR;ref=1;x=350;z=120",
        "  [Schlitten fährt auf Anfahrposition, Kreisbogen, Rückfahrt]",
        "← EVT;7;OK;DOOR_OPEN_DONE;x=350;z=120",
        "← EVT;0;STATE;READY;ref=1;x=350;z=120",
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    add_heading(doc, "Plattenentnahme (PICKUP)", level=3)
    for line in [
        "→ CMD;5;PICKUP;gripper_depth=120;lift_offset=8",
        "← RSP;5;ACK",
        "← EVT;0;STATE;BUSY_PICKUP;ref=1;x=350;z=120",
        "← EVT;5;OK;PICKUP_DONE;x=350;z=128",
        "← EVT;0;STATE;READY;ref=1;x=350;z=128",
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    add_heading(doc, "Fehlerfall und Quittierung", level=3)
    for line in [
        "→ CMD;4;MOVE_TO;x=500;z=120",
        "← RSP;4;ACK",
        "← EVT;0;STATE;BUSY_MOVING;ref=1;x=350;z=120",
        "← EVT;0;ERR;OBSTACLE;x=412;z=120",
        "← EVT;0;STATE;ERROR;ref=1;x=412;z=120",
        "← EVT;0;STATUS;state=ERROR;error=OBSTACLE;...",
        "→ CMD;5;RESET_ERROR",
        "← RSP;5;ACK",
        "← EVT;5;OK;ERROR_RESET",
        "← EVT;0;STATE;NOT_REFERENCED;ref=0;x=0;z=0",
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    add_divider(doc)

    # 1.7 Plattenwechsel-Sequenz
    add_heading(doc, "1.7  Plattenwechsel-Sequenz (Pi-Logik)", level=2)
    add_para(doc, "Der Pi orchestriert alle Schritte über die o.g. Befehle:")
    for line in [
        "Schritt  1: MOVE_TO  (drucker.pos_x, drucker.pos_z_anfahr)",
        "Schritt  2: OPEN_DOOR (x_approach=pos_x, z_approach=pos_z_tuer,",
        "                       arm_extend, radius, angle)",
        "Schritt  3: STATUS → door_open == 1? (sonst TUERFEHLER)",
        "Schritt  4: MOVE_TO  (drucker.pos_x, drucker.pos_z_druckbett)",
        "Schritt  5: PICKUP   (gripper_depth, lift_offset)",
        "Schritt  6: MOVE_TO  (drucker.pos_x, drucker.pos_z_anfahr)",
        "Schritt  7: CLOSE_DOOR (x_approach=x_close, ...)",
        "Schritt  8: MOVE_TO  (ablage.x, ablage.z)",
        "Schritt  9: DEPOSIT  (gripper_depth, lift_offset)   ← alte Platte abgelegt",
        "Schritt 10: MOVE_TO  (magazin.x, magazin.z)",
        "Schritt 11: PICKUP   (gripper_depth, lift_offset)   ← neue Platte",
        "Schritt 12: MOVE_TO  (drucker.pos_x, drucker.pos_z_anfahr)",
        "Schritt 13: OPEN_DOOR (...)",
        "Schritt 14: STATUS → door_open == 1?",
        "Schritt 15: MOVE_TO  (drucker.pos_x, drucker.pos_z_druckbett)",
        "Schritt 16: DEPOSIT  (gripper_depth, lift_offset)   ← neue Platte eingesetzt",
        "Schritt 17: MOVE_TO  (drucker.pos_x, drucker.pos_z_anfahr)",
        "            CLOSE_DOOR (...)",
        "            MOVE_HOME",
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # 2. MQTT
    # ════════════════════════════════════════════════════════════════
    add_heading(doc, "2  Pi ↔ Außenwelt: MQTT")
    add_para(doc,
        "Broker: Mosquitto (lokal auf dem Pi, localhost:1883). Alle Topics verwenden den "
        "Base-Topic aus config.yaml (Standard: plattenwechsler-g6twie23a).")

    add_heading(doc, "2.1  Status-Topics (retained)", level=2)
    add_table(doc,
        ["Topic", "Inhalt", "Trigger"],
        [
            ["<base>/status/system",        "Vollständiger Pi-Snapshot (JSON)",    "jede Zustandsänderung"],
            ["<base>/status/esp",           "ESP-Status (JSON)",                   "jedes ESP-STATUS-Event"],
            ["<base>/status/online",        "online / offline",                    "Verbindung / Will-Message"],
            ["<base>/status/drucker/<id>",  "Drucker-Status (JSON)",               "Fertig-Pin, Zustandsänderung"],
        ],
        col_widths=[5.5, 5.5, 5.5]
    )

    add_heading(doc, "2.2  Event-Topics", level=2)
    add_table(doc,
        ["Topic", "Inhalt", "Trigger"],
        [
            ["<base>/error",          "Fehlerdetails (JSON)",             "neuer Fehler"],
            ["<base>/event/auftrag",  "Ergebnis Plattenwechsel (JSON)",   "abgeschlossener Auftrag"],
        ],
        col_widths=[5.5, 5.5, 5.5]
    )

    add_heading(doc, "2.3  Befehls-Topics", level=2)
    add_para(doc,
        "Alle Befehle werden als JSON-Payload gesendet. Ergebnisse kommen über "
        "die Status-Topics zurück.")
    add_table(doc,
        ["Topic", "Payload", "Wirkung"],
        [
            ["<base>/cmd/auftrag",          '{"drucker_id": N}',                     "Plattenwechsel starten"],
            ["<base>/cmd/quittieren",       "{}",                                    "Fehler quittieren"],
            ["<base>/cmd/referenzfahrt",    "{}",                                    "Referenzfahrt starten"],
            ["<base>/cmd/stop",             "{}",                                    "Sofort-Stop"],
            ["<base>/cmd/service/modus",    '{"aktiv": true|false}',                 "Service-Modus"],
            ["<base>/cmd/service/fahre",    '{"ziel": "home"|"ablage"|"magazin"}',   "Service-Fahrt"],
            ["<base>/cmd/service/drucker",  '{"drucker_id": N}',                     "Zu Drucker fahren (Service)"],
            ["<base>/cmd/drucker/setzen",   "Drucker-Objekt (s.u.)",                 "Drucker anlegen/aktualisieren"],
            ["<base>/cmd/drucker/entfernen",'{"id": N}',                             "Drucker entfernen"],
        ],
        col_widths=[5.5, 4.5, 6.5]
    )

    add_heading(doc, "Drucker-Objekt für drucker/setzen:", level=3)
    for line in [
        '{',
        '  "id": 1,  "name": "Drucker 1",  "pin_fertig": 27,',
        '  "pos_x": 370,  "pos_z_anfahr": 120,  "pos_z_tuer": 105,',
        '  "pos_z_druckbett": 50,  "door_arm_hub_mm": 30,',
        '  "tuer_radius": 150,  "tuer_winkel": 160,',
        '  "gripper_depth": 120,  "lift_offset": 8',
        '}',
    ]:
        add_mono(doc, line)
    doc.add_paragraph()

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # 3. Web-App API
    # ════════════════════════════════════════════════════════════════
    add_heading(doc, "3  Pi ↔ Browser: Web-App API")
    add_para(doc, "Die Flask-Web-App läuft auf Port 5000.")

    add_heading(doc, "3.1  GET-Endpunkte", level=2)
    add_table(doc,
        ["Endpunkt", "Rückgabe"],
        [
            ["GET /",            "HTML-Oberfläche (Single Page)"],
            ["GET /api/state",   "Aktueller System-Snapshot (JSON)"],
            ["GET /api/events",  "Server-Sent Events (SSE-Stream), Content-Type: text/event-stream"],
        ],
        col_widths=[5.0, 11.5]
    )
    add_para(doc, "SSE-Events:", bold=True)
    add_bullet(doc, "state — System-Snapshot bei jeder Zustandsänderung")
    add_bullet(doc, "error — Fehlerdetails")
    doc.add_paragraph()

    add_heading(doc, "3.2  POST-Endpunkte", level=2)
    add_para(doc,
        "Alle POST-Endpunkte erwarten JSON-Body und geben "
        '{"ok": true} oder {"error": "..."} zurück.')
    add_table(doc,
        ["Endpunkt", "Body", "Wirkung"],
        [
            ["POST /api/cmd/auftrag",           '{"drucker_id": N}',                   "Plattenwechsel starten"],
            ["POST /api/cmd/quittieren",        "{}",                                  "Fehler quittieren"],
            ["POST /api/cmd/referenzfahrt",     "{}",                                  "Referenzfahrt starten"],
            ["POST /api/cmd/stop",              "{}",                                  "Sofort-Stop"],
            ["POST /api/cmd/service/modus",     '{"aktiv": true|false}',               "Service-Modus"],
            ["POST /api/cmd/service/fahre",     '{"ziel": "home"|"ablage"|"magazin"}', "Service-Fahrt"],
            ["POST /api/cmd/service/drucker",   '{"drucker_id": N}',                   "Zu Drucker fahren"],
            ["POST /api/cmd/drucker/setzen",    "Drucker-Objekt",                      "Drucker anlegen/aktualisieren"],
            ["POST /api/cmd/drucker/entfernen", '{"id": N}',                           "Drucker entfernen"],
        ],
        col_widths=[5.5, 4.5, 6.5]
    )

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # 4. GPIO
    # ════════════════════════════════════════════════════════════════
    add_heading(doc, "4  Pi: GPIO-Belegung")
    add_para(doc,
        "Alle Pins in BCM-Nummerierung. Konfigurierbar in config.yaml, Abschnitt gpio.")

    add_heading(doc, "4.1  Eingänge (Pi ← Hardware)", level=2)
    add_table(doc,
        ["Pin (BCM)", "Funktion", "Auslöser", "Pi-Reaktion"],
        [
            ["17 (Standard)", "Endschalter X-Achse", "Schlitten an Home-Position X",
             "Sofort HOME_SWITCH_HIT;axis=X an ESP"],
            ["4 (Standard)",  "Endschalter Z-Achse", "Schlitten an Home-Position Z",
             "Sofort HOME_SWITCH_HIT;axis=Z an ESP"],
            ["26 (Standard)", "Not-Aus",             "Sicherheitsschalter",
             "STOP an ESP, Zustand → FEHLER"],
            ["pro Drucker",   "Fertig-Pin",          "Drucker meldet Druck fertig (LOW)",
             "Auftrag in Queue einreihen"],
        ],
        col_widths=[2.8, 3.5, 4.2, 6.0]
    )
    add_note(doc,
        "Die Endschalter-Pins und der Not-Aus-Pin sind in config.yaml konfigurierbar. "
        "Fertig-Pins werden pro Drucker vergeben und müssen eindeutig sein (keine Doppelbelegung).")

    add_heading(doc, "4.2  Signalverarbeitung", level=2)
    add_table(doc,
        ["Signal", "Verarbeitung"],
        [
            ["Endschalter X/Z", "Pi sendet sofort HOME_SWITCH_HIT;axis=<X|Z> an ESP"],
            ["Not-Aus",         "Pi sendet STOP an ESP; Zustand wechselt zu FEHLER / NOT_AUS"],
            ["Drucker Fertig",  "Pi legt Auftrag für diesen Drucker in die Queue"],
        ],
        col_widths=[4.0, 12.5]
    )

    add_heading(doc, "4.3  UART-Verbindung zum ESP", level=2)
    add_table(doc,
        ["Parameter", "Wert"],
        [
            ["Port",      "/dev/ttyUSB0  (konfigurierbar)"],
            ["Baudrate",  "115200 Baud, 8N1"],
            ["Richtung",  "bidirektional, vollduplex"],
        ],
        col_widths=[4.0, 12.5]
    )

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════
    # 5. config.yaml
    # ════════════════════════════════════════════════════════════════
    add_heading(doc, "5  Konfigurationsschnittstelle (config.yaml)")
    add_para(doc,
        "Die Konfiguration wird beim Start geladen und bei Änderungen über die HMI oder "
        "MQTT direkt zurückgeschrieben.")

    add_heading(doc, "5.1  Abschnitt system", level=2)
    add_table(doc,
        ["Schlüssel", "Standardwert", "Bedeutung"],
        [
            ["log_level", "INFO",  "Logging-Level: DEBUG | INFO | WARNING | ERROR"],
            ["log_file",  "(Pfad)", "Pfad zur Log-Datei"],
        ],
        col_widths=[4.0, 4.0, 8.5]
    )

    add_heading(doc, "5.2  Abschnitt esp", level=2)
    add_table(doc,
        ["Schlüssel", "Standardwert", "Bedeutung"],
        [
            ["port",                  "/dev/ttyUSB0", "Serieller Port"],
            ["baud",                  "115200",        "Baudrate"],
            ["ack_timeout_s",         "1.0",           "Max. Wartezeit auf RSP;ACK"],
            ["move_timeout_s",        "60.0",          "Timeout für MOVE_TO, OPEN/CLOSE_DOOR"],
            ["home_timeout_s",        "30.0",          "Timeout für HOME, MOVE_HOME"],
            ["mech_timeout_s",        "10.0",          "Timeout für PICKUP, DEPOSIT"],
            ["heartbeat_timeout_s",   "5.0",           "Reconnect nach dieser Zeit ohne HEARTBEAT"],
            ["reconnect_intervall_s", "2.0",           "Wartezeit zwischen Reconnect-Versuchen"],
            ["skip_homing",           "false",         "Referenzfahrt überspringen (nur für Tests)"],
        ],
        col_widths=[4.5, 3.0, 9.0]
    )

    add_heading(doc, "5.3  Abschnitt gpio", level=2)
    add_table(doc,
        ["Schlüssel", "Bedeutung"],
        [
            ["gpio.endschalter.x",  "BCM-Pin Endschalter X-Achse (Standard: 17)"],
            ["gpio.endschalter.z",  "BCM-Pin Endschalter Z-Achse (Standard: 4)"],
            ["gpio.not_aus.pin",    "BCM-Pin Not-Aus (Standard: 26)"],
            ["gpio.not_aus.invert", "true = active-high (Standard: false)"],
        ],
        col_widths=[4.5, 12.0]
    )

    add_heading(doc, "5.4  Abschnitt drucker (pro Drucker)", level=2)
    add_table(doc,
        ["Schlüssel", "Typ", "Bedeutung"],
        [
            ["id",              "int",      "Eindeutige Drucker-ID (1–n)"],
            ["name",            "string",   "Anzeigename in der HMI"],
            ["pin_fertig",      "int (BCM)","GPIO-Pin Fertig-Signal (0 = deaktiviert, muss eindeutig sein)"],
            ["pos_x",           "int (mm)", "X-Position vor dem Drucker"],
            ["pos_z_anfahr",    "int (mm)", "Z-Höhe Ausgangsposition"],
            ["pos_z_tuer",      "int (mm)", "Z-Höhe für OPEN/CLOSE_DOOR (z_approach)"],
            ["pos_z_druckbett", "int (mm)", "Z-Höhe Gabel-Bereitschaftsposition"],
            ["door_arm_hub_mm", "int (mm)", "Türarm-Ausfahrlänge (arm_extend)"],
            ["tuer_radius",     "int (mm)", "Kreisradius Türbewegung"],
            ["tuer_winkel",     "int (°)",  "Öffnungswinkel Tür"],
            ["gripper_depth",   "int (mm)", "Greifer-Ausfahrtiefe"],
            ["lift_offset",     "int (mm)", "Z-Hub bei PICKUP/DEPOSIT"],
        ],
        col_widths=[3.8, 2.2, 10.5]
    )

    add_heading(doc, "5.5  Abschnitt ablagen", level=2)
    add_table(doc,
        ["Schlüssel", "Typ", "Bedeutung"],
        [
            ["id, name",      "int, string", "Eindeutige ID, Anzeigename"],
            ["x, z",          "int (mm)",    "Position"],
            ["gripper_depth", "int (mm)",    "Greifer-Ausfahrtiefe"],
            ["lift_offset",   "int (mm)",    "Z-Hub"],
            ["belegt",        "bool",        "true = Platte liegt auf Ablage (Laufzeit-Status)"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )

    add_heading(doc, "5.6  Abschnitt magazine", level=2)
    add_table(doc,
        ["Schlüssel", "Typ", "Bedeutung"],
        [
            ["id, name",      "int, string", "Eindeutige ID, Anzeigename"],
            ["x, z",          "int (mm)",    "Position"],
            ["gripper_depth", "int (mm)",    "Greifer-Ausfahrtiefe"],
            ["lift_offset",   "int (mm)",    "Z-Hub"],
            ["verfuegbar",    "bool",        "true = Platte vorhanden (Laufzeit-Status)"],
        ],
        col_widths=[3.5, 2.5, 10.5]
    )

    add_heading(doc, "5.7  Abschnitt mqtt", level=2)
    add_table(doc,
        ["Schlüssel", "Standardwert", "Bedeutung"],
        [
            ["enabled",               "true",                      "MQTT aktivieren"],
            ["broker_host",           "localhost",                 "Broker-Adresse"],
            ["broker_port",           "1883",                      "Broker-Port"],
            ["client_id",             "plattenwechsler-pi",        "MQTT Client-ID"],
            ["base_topic",            "plattenwechsler-g6twie23a", "Basis-Topic"],
            ["qos",                   "1",                         "Quality of Service"],
            ["retain_status",         "true",                      "Status retained veröffentlichen"],
            ["reconnect_intervall_s", "5.0",                       "Reconnect-Intervall"],
        ],
        col_widths=[4.5, 4.0, 8.0]
    )

    add_heading(doc, "5.8  Abschnitt telegram", level=2)
    add_table(doc,
        ["Schlüssel", "Bedeutung"],
        [
            ["enabled",              "Telegram-Bot aktivieren"],
            ["bot_token",            "Bot-Token von @BotFather"],
            ["allowed_chat_ids",     "Liste autorisierter Chat-IDs"],
            ["send_status_to_first", "Status-Updates nach Plattenwechsel/Fehler an ersten Chat"],
        ],
        col_widths=[4.5, 12.0]
    )

    add_heading(doc, "5.9  Abschnitt hmi", level=2)
    add_table(doc,
        ["Schlüssel", "Standardwert", "Bedeutung"],
        [
            ["fullscreen", "true", "Vollbild-Modus"],
            ["width",      "1280", "Fenstbreite (Pixel)"],
            ["height",     "720",  "Fensterhöhe (Pixel)"],
        ],
        col_widths=[3.5, 3.5, 9.5]
    )

    doc.add_page_break()

    # Anhang Timing
    add_heading(doc, "Anhang: Timing-Parameter")
    add_table(doc,
        ["Parameter", "Wert", "Konfigurierbar in"],
        [
            ["UART Baudrate",                          "115200 Baud",  "config.yaml esp.baud"],
            ["Heartbeat-Intervall (ESP → Pi)",         "1000 ms",      "ESP-Firmware (fest)"],
            ["Heartbeat-Timeout (Pi)",                 "5000 ms",      "esp.heartbeat_timeout_s"],
            ["STATUS-Stream-Intervall",                "100 ms",       "ESP-Firmware (fest)"],
            ["Hindernissensor-Abfrageintervall",       "50 ms",        "ESP-Firmware (fest)"],
            ["ACK-Timeout (Pi wartet auf RSP)",        "1000 ms",      "esp.ack_timeout_s"],
            ["Timeout MOVE_TO / OPEN/CLOSE_DOOR",      "60.000 ms",    "esp.move_timeout_s"],
            ["Timeout HOME / MOVE_HOME",               "30.000 ms",    "esp.home_timeout_s"],
            ["Timeout PICKUP / DEPOSIT",               "10.000 ms",    "esp.mech_timeout_s"],
            ["MQTT Reconnect-Intervall",               "5000 ms",      "mqtt.reconnect_intervall_s"],
            ["ESP Reconnect-Intervall",                "2000 ms",      "esp.reconnect_intervall_s"],
        ],
        col_widths=[7.0, 3.0, 6.5]
    )

    out = "/home/nura/plattenwechsler/docs/Schnittstellen_Doku.docx"
    doc.save(out)
    print(f"Gespeichert: {out}")


if __name__ == "__main__":
    build()
